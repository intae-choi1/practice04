from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart
wb = load_workbook("sample.xlsx")
ws = wb.active

# 데이터 위치 지정
bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
bar_chart = BarChart() # 차트 종류 설정
bar_chart.add_data(bar_value) # 차트 데이터 추가

ws.add_chart(bar_chart, "E1") # 차트 넣을 위치

line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data=True) # 계열이 데이터로부터 지정됨
line_chart.style = 20 # 미리 지정된 스타일 적용 (사용자 지정도 가능)
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"

ws.add_chart(line_chart, "E14")

wb.save("sample_chart.xlsx")